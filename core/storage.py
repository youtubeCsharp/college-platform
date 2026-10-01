# -*- coding: utf-8 -*-
"""Excel 结构化存储引擎。

设计要点(对应可行性方案):
- 三办 Excel 物理隔离,互不冲突;
- 单进程内按文件加线程锁,串行化写入;
- 保存采用"临时文件 + 原子替换",杜绝写坏文件;
- 文件被 Excel 直接打开占用时,抛出友好提示而非静默失败;
- 操作日志追加写入 jsonl,只增不删,保证可追溯。
"""
import json
import os
import threading
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent.parent
DATA_DIR = BASE_DIR / "data"
LEDGER_DIR = DATA_DIR / "台账"
TASK_DIR = DATA_DIR / "任务"
LOG_DIR = DATA_DIR / "日志"
FILES_DIR = BASE_DIR / "files"
BACKUP_DIR = BASE_DIR / "backups"
LOG_FILE = LOG_DIR / "操作日志.jsonl"

OFFICES = ["综合办", "教学办", "学生办"]
OFFICE_PREFIX = {"综合办": "ZH", "教学办": "JX", "学生办": "XS"}

LEDGER_COLUMNS = ["编号", "日期", "工作大类", "工作内容", "完成情况", "参与人员",
                  "备注", "敏感事项", "附件", "录入人", "录入时间",
                  "最后修改人", "最后修改时间", "状态"]

TASK_COLUMNS = ["任务编号", "任务标题", "指派部门", "责任人", "截止日期", "优先级",
                "工作说明", "状态", "需协调资源", "发布人", "发布时间", "办结时间"]
REPLY_COLUMNS = ["回复编号", "任务编号", "回复人", "回复时间", "回复内容", "进度状态", "申请协调"]

HEADER_FILL = PatternFill("solid", fgColor="1F3A6E")
HEADER_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

LEDGER_WIDTHS = [17, 12, 12, 48, 10, 16, 20, 9, 30, 10, 19, 10, 19, 8]
TASK_WIDTHS = [17, 32, 11, 10, 12, 8, 42, 9, 12, 10, 19, 19]
REPLY_WIDTHS = [22, 17, 10, 19, 50, 10, 10]

_locks = {}
_locks_guard = threading.Lock()


class BusyFileError(Exception):
    """Excel 文件正被其他程序(如 Microsoft Excel)占用。"""


def ledger_path(office: str) -> Path:
    return LEDGER_DIR / f"{office}.xlsx"


def task_path() -> Path:
    return TASK_DIR / "工作安排.xlsx"


def _lock_for(path: Path) -> threading.RLock:
    with _locks_guard:
        key = str(path)
        if key not in _locks:
            _locks[key] = threading.RLock()
        return _locks[key]


def _style_sheet(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
    ws.freeze_panes = "A2"


def _create_file(path: Path, sheets: dict):
    """sheets: {工作表名: (列名列表, 列宽列表)}"""
    wb = Workbook()
    wb.remove(wb.active)
    for name, (cols, widths) in sheets.items():
        ws = wb.create_sheet(name)
        ws.append(cols)
        _style_sheet(ws, widths)
    tmp = path.with_name(path.stem + ".tmp.xlsx")
    wb.save(tmp)
    wb.close()
    os.replace(tmp, path)


def init_storage():
    for d in (LEDGER_DIR, TASK_DIR, LOG_DIR, FILES_DIR, BACKUP_DIR):
        d.mkdir(parents=True, exist_ok=True)
    for office in OFFICES:
        p = ledger_path(office)
        if not p.exists():
            _create_file(p, {"台账": (LEDGER_COLUMNS, LEDGER_WIDTHS)})
    p = task_path()
    if not p.exists():
        _create_file(p, {"任务表": (TASK_COLUMNS, TASK_WIDTHS),
                         "回复表": (REPLY_COLUMNS, REPLY_WIDTHS)})


def read_rows(path: Path, sheet: str) -> list:
    with _lock_for(path):
        try:
            wb = load_workbook(path, data_only=True)
        except PermissionError:
            raise BusyFileError(f"{path.name} 正被其他程序占用(可能有人用 Excel 直接打开了文件),请关闭后重试。")
        try:
            ws = wb[sheet]
            headers = [c.value for c in ws[1]]
            rows = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row is None or all(v is None or v == "" for v in row):
                    continue
                item = {}
                for i, h in enumerate(headers):
                    v = row[i] if i < len(row) else None
                    item[h] = "" if v is None else (v if isinstance(v, str) else str(v))
                rows.append(item)
            return rows
        finally:
            wb.close()


def write_rows(path: Path, sheet: str, rows: list):
    """整体重写指定工作表(先写临时文件,再原子替换)。"""
    with _lock_for(path):
        try:
            wb = load_workbook(path)
        except PermissionError:
            raise BusyFileError(f"{path.name} 正被其他程序占用(可能有人用 Excel 直接打开了文件),请关闭后重试。")
        tmp = path.with_name(path.stem + ".tmp.xlsx")
        try:
            ws = wb[sheet]
            headers = [c.value for c in ws[1]]
            if ws.max_row and ws.max_row > 1:
                ws.delete_rows(2, ws.max_row - 1)
            for r in rows:
                values = []
                for h in headers:
                    v = r.get(h, "")
                    values.append("" if v is None else (v if isinstance(v, str) else str(v)))
                ws.append(values)
            try:
                wb.save(tmp)
            except PermissionError:
                raise BusyFileError(f"{path.name} 正被其他程序占用,无法保存,请关闭 Excel 后重试。")
            wb.close()
            os.replace(tmp, path)
        finally:
            try:
                wb.close()
            except Exception:
                pass
            if tmp.exists():
                try:
                    tmp.unlink()
                except OSError:
                    pass


def append_log(user: str, action: str, detail: str = ""):
    entry = {"时间": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
             "用户": user, "操作": action, "详情": detail}
    with _lock_for(LOG_FILE):
        with open(LOG_FILE, "a", encoding="utf-8") as f:
            f.write(json.dumps(entry, ensure_ascii=False) + "\n")
